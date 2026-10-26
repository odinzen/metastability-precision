"""Step 1 coverage count: which automated-synthesis targets and precursors have a measured
formation enthalpy, and how many measured compounds sit in each target's chemical system.

Measured set: Kingsbury/Wang et al. compilation (Kim 2017, Kubaschewski, NIST-JANAF), 2135 solids.
Matching is by reduced formula only, so a polymorph mismatch is possible and is flagged later.
"""
import gzip
import json
import re
from itertools import combinations
from pathlib import Path

import openpyxl
import pandas as pd
from pymatgen.core import Composition

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "results"


def reduced(formula: str) -> str:
    return Composition(formula).reduced_formula


def load_measured() -> pd.DataFrame:
    d = json.load(gzip.open(RAW / "measured" / "expt_formation_enthalpy_kingsbury.json.gz"))
    df = pd.DataFrame(d["data"], columns=d["columns"])
    df["reduced"] = df["formula"].map(reduced)
    df["elements"] = df["formula"].map(lambda f: frozenset(str(e) for e in Composition(f).elements))
    return df


def chen_sets():
    wb = openpyxl.load_workbook(RAW / "chen2024" / "si" / "MOESM9_ESM.xlsx", read_only=True, data_only=True)
    rows = list(wb["Figure 4D"].iter_rows(values_only=True))[1:]
    targets = {r[1]: dict(inverse_hull_eV=float(r[2]), reaction_eV=float(r[3]), score_diff=float(r[4])) for r in rows}
    wb = openpyxl.load_workbook(RAW / "chen2024" / "si" / "MOESM7_ESM.xlsx", read_only=True, data_only=True)
    precursors = {}
    for ws in wb.worksheets:
        hdr = next(ws.iter_rows(values_only=True))
        sets = {re.search(r"\[(.*?)\]", str(h)).group(1) for h in hdr if h and "[" in str(h)}
        precursors[ws.title] = sorted({p.strip(" '") for s in sets for p in s.split(",")})
    return targets, precursors


def alab_sets():
    df = pd.read_csv(RAW / "alab2023" / "unzipped" / "20230502 Synthesis Results with Recipes.csv", encoding="latin-1")
    prec = {t: sorted(set(re.findall(r"mg\s+([A-Za-z0-9()]+)", s))) for t, s in zip(df["Target"], df["Precursors + Masses"].fillna(""))}
    return dict(zip(df["Target"], df["Result"])), prec


def system_of(formula: str) -> frozenset:
    return frozenset(str(e) for e in Composition(formula).elements)


def main():
    OUT.mkdir(exist_ok=True)
    meas = load_measured()
    measured = set(meas["reduced"])
    chen_t, chen_p = chen_sets()
    alab_t, alab_p = alab_sets()

    rows = []
    for name, info in chen_t.items():
        rows.append(dict(dataset="Chen2024", role="target", formula=name, outcome=info["score_diff"],
                         inverse_hull_meV=round(1000 * info["inverse_hull_eV"], 2)))
    for t, ps in chen_p.items():
        rows += [dict(dataset="Chen2024", role="precursor", formula=p, target=t) for p in ps]
    for name, res in alab_t.items():
        rows.append(dict(dataset="A-Lab2023", role="target", formula=name, outcome=res))
    for t, ps in alab_p.items():
        rows += [dict(dataset="A-Lab2023", role="precursor", formula=p, target=t) for p in ps]
    df = pd.DataFrame(rows)
    df["reduced"] = df["formula"].map(reduced)
    df["measured"] = df["reduced"].isin(measured)
    df.to_csv(OUT / "coverage_compounds.csv", index=False)

    # Measured compounds inside each target's chemical system (all sub-systems, elements excluded).
    sysrows = []
    for ds, targets in (("Chen2024", chen_t), ("A-Lab2023", alab_t)):
        for t in targets:
            els = system_of(t)
            inside = meas[meas["elements"].map(lambda e: e <= els and len(e) > 1)]
            n_ternary_up = int((inside["elements"].map(len) >= 3).sum())
            sysrows.append(dict(dataset=ds, target=t, system="-".join(sorted(els)), n_elements=len(els),
                                measured_in_system=len(inside), measured_ternary_or_higher=n_ternary_up,
                                target_measured=reduced(t) in measured))
    sysdf = pd.DataFrame(sysrows)
    sysdf.to_csv(OUT / "coverage_systems.csv", index=False)

    lines = []
    for ds in ("Chen2024", "A-Lab2023"):
        d = df[df.dataset == ds]
        tg = d[d.role == "target"]
        pr = d[d.role == "precursor"].drop_duplicates("reduced")
        s = sysdf[sysdf.dataset == ds]
        lines.append(f"{ds}: targets measured {tg.measured.sum()}/{len(tg)}; unique precursors measured "
                     f"{pr.measured.sum()}/{len(pr)} (missing: {', '.join(sorted(pr[~pr.measured].formula))}); "
                     f"median measured compounds per target system {s.measured_in_system.median():.0f}, "
                     f"median ternary-or-higher {s.measured_ternary_or_higher.median():.0f}")
    print("\n".join(lines))
    print("\nTargets with a measured value:")
    print(df[(df.role == "target") & df.measured][["dataset", "formula", "outcome"]].to_string(index=False))
    print("\nPer-system detail (sorted by measured ternary-or-higher):")
    print(sysdf.sort_values("measured_ternary_or_higher", ascending=False).to_string(index=False))


if __name__ == "__main__":
    main()
