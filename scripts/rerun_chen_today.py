"""Rerun Chen et al.'s precursor-selection code (their repo, commit 6dcee76) on today's Materials Project
entries, and on their own cached entries where the repo ships them, to see whether the decisions move.

Their getOrigStableEntriesList() is swapped for a loader that reads our cache, so nothing calls an API and
nothing is written into their folder.
"""
import gzip
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
import pandas as pd
from monty.json import MontyDecoder
from pymatgen.analysis.phase_diagram import PhaseDiagram
from pymatgen.core import Composition
from pymatgen.entries.computed_entries import ComputedEntry

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]
CODE = ROOT / "data" / "raw" / "chen2024" / "code" / "synthesis_planning_algorithm"
sys.path.insert(0, str(CODE))
import synthesis_planning.synthesis_pathways as sp_mod  # noqa: E402

TRADITIONAL = re.compile(r"CO3|NH4")


def today_stable(els):
    chemsys = "-".join(sorted(els))
    path = ROOT / "data" / "raw" / "mp" / f"{chemsys}.json.gz"
    entries = MontyDecoder().process_decoded(json.load(gzip.open(path)))
    # Plain ComputedEntry, as their cache holds, so their name-based matching behaves the same.
    plain = [ComputedEntry(e.composition, e.energy, correction=0, entry_id=str(i)) for i, e in enumerate(entries)]
    return list(PhaseDiagram(plain).stable_entries)


def chen_cached(els):
    path = CODE / "data" / "entries_files" / "-".join(sorted(els))
    return [ComputedEntry.from_dict(d) for d in json.load(open(path))]


def chen_published():
    wb = openpyxl.load_workbook(ROOT / "data/raw/chen2024/si/MOESM9_ESM.xlsx", read_only=True, data_only=True)
    pub = {r[1]: dict(pub_invE_meV=1000 * float(r[2]), pub_rxn_meV=1000 * float(r[3]), score_diff=float(r[4]))
           for r in list(wb["Figure 4D"].iter_rows(values_only=True))[1:]}
    wb = openpyxl.load_workbook(ROOT / "data/raw/chen2024/si/MOESM7_ESM.xlsx", read_only=True, data_only=True)
    pairs = {}
    for ws in wb.worksheets:
        sets = {tuple(sorted(s.replace("'", "").replace(" ", "").split(",")))
                for h in next(ws.iter_rows(values_only=True)) if h and "[" in str(h)
                for s in [re.search(r"\[(.*?)\]", str(h)).group(1)]}
        # Their method selects precursor pairs; a three-precursor set with the pair inside is a variant.
        robot = [s for s in sets if len(s) == 2 and not any(TRADITIONAL.search(p) for p in s)]
        pairs[norm(ws.title)] = robot
    for t in pub:
        pub[t]["robot_pairs"] = pairs.get(norm(t), [])
    return pub


def norm(f):
    # Chen's own normalisation; reduced_formula leaves Li0.5Nb0.5W0.5O3 fractional.
    return Composition(f).get_integer_formula_and_factor()[0]


def run(target, loader):
    sp_mod.getOrigStableEntriesList = loader
    sp = sp_mod.SynthesisPathways(target, exclude_reactants=["O2"], selected_reactions_to_csv=False)
    ranked = [(tuple(sorted(norm(e.name) for e in r.reactants)), 1000 * r.reactE, 1000 * r.invE)
              for r in sp.selected_reactions]
    allp = {tuple(sorted(norm(e.name) for e in r.reactants)): (1000 * r.reactE, 1000 * r.invE)
            for r in sp.all_pairwise_reactions}
    return ranked, allp


def main():
    pub = chen_published()
    rows = []
    for target, p in pub.items():
        robot = [tuple(sorted(norm(x) for x in s)) for s in p["robot_pairs"]]
        robot_pair = robot[0] if len(robot) == 1 else None
        row = dict(target=target, robot_pair="+".join(robot_pair) if robot_pair else ";".join("+".join(r) for r in robot),
                   pub_rxn_meV=round(p["pub_rxn_meV"], 1), pub_invE_meV=round(p["pub_invE_meV"], 1),
                   score_diff=p["score_diff"])
        for label, loader in (("today", today_stable), ("chen_cache", chen_cached)):
            try:
                ranked, allp = run(target, loader)
            except FileNotFoundError:
                continue
            top = ranked[0] if ranked else None
            row[f"{label}_top_pair"] = "+".join(top[0]) if top else "none selected"
            row[f"{label}_robot_pair_still_selected"] = bool(robot_pair) and any(r[0] == robot_pair for r in ranked)
            row[f"{label}_robot_pair_is_top"] = bool(top) and top[0] == robot_pair
            if robot_pair and robot_pair in allp:
                row[f"{label}_rxn_meV"] = round(allp[robot_pair][0], 1)
                row[f"{label}_invE_meV"] = round(allp[robot_pair][1], 1)
        rows.append(row)
        print(target, {k: v for k, v in row.items() if k.startswith(("today", "chen_cache"))}, flush=True)
    out = pd.DataFrame(rows)
    out.to_csv(ROOT / "results" / "chen_rerun_today.csv", index=False)


if __name__ == "__main__":
    main()
